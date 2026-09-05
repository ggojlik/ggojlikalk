import sys
import os
import subprocess

import pandas as pd

from PyQt6.QtCore import Qt
from PyQt6.QtWidgets import (
    QApplication,
    QWidget,
    QVBoxLayout,
    QHBoxLayout,
    QComboBox,
    QPushButton,
    QTableWidget,
    QTableWidgetItem,
    QLabel,
    QSlider,
    QGroupBox,
    QMessageBox
)

from openpyxl import Workbook

from matplotlib.backends.backend_qtagg import FigureCanvasQTAgg as FigureCanvas
from matplotlib.figure import Figure


# ==========================================
# ŚCIEŻKI DO PLIKÓW
# ==========================================

KATALOG_PROGRAMU = os.path.dirname(
    os.path.abspath(__file__)
)

# Oryginalny plik z danymi
SCIEZKA_PLIKU = os.path.join(
    KATALOG_PROGRAMU,
    "dane_50_fikcyjnych_pacjentow.xlsx"
)

# NOWY plik z wynikami
SCIEZKA_WYNIKI = os.path.join(
    KATALOG_PROGRAMU,
    "wyniki_gui.xlsx"
)


# ==========================================
# WCZYTANIE DANYCH
# ==========================================

dane = pd.read_excel(
    SCIEZKA_PLIKU
)


# ==========================================
# KONWERSJA BMI I TEMPERATURY NA LICZBY
# ==========================================

dane["BMI"] = pd.to_numeric(
    dane["BMI"],
    errors="coerce"
)

dane["Temperatura (°C)"] = pd.to_numeric(
    dane["Temperatura (°C)"],
    errors="coerce"
)


# ==========================================
# KONWERSJA WIEKU NA LICZBY
# ==========================================

dane["Wiek"] = pd.to_numeric(
    dane["Wiek"],
    errors="coerce"
)


# ==========================================
# GRUPY WIEKOWE
# ==========================================

dane["Grupa wiekowa"] = pd.cut(
    dane["Wiek"],
    bins=[0, 18, 40, 60, 80, 120],
    labels=[
        "0-18",
        "19-40",
        "41-60",
        "61-80",
        "81+"
    ],
    include_lowest=True
)


# ==========================================
# PACJENCI Z GORĄCZKĄ
# ==========================================

goraczka = dane[
    dane["Gorączka"]
    .astype(str)
    .str.lower()
    .str.strip() == "tak"
].copy()


# ==========================================
# UJEDNOLICENIE TAK/NIE
# ==========================================

goraczka["Czy brał(a) leki"] = (
    goraczka["Czy brał(a) leki"]
    .astype(str)
    .str.strip()
    .str.lower()
    .map({
        "tak": "Tak",
        "nie": "Nie"
    })
)


# ==========================================
# GUI
# ==========================================

class Okno(QWidget):

    def __init__(self):

        super().__init__()

        self.setWindowTitle(
            "System analizy pacjentów z gorączką"
        )

        self.resize(
            1400,
            900
        )

        layout = QVBoxLayout()


        # ======================================
        # TYTUŁ
        # ======================================

        tytul = QLabel(
            "ANALIZA PACJENTÓW Z GORĄCZKĄ"
        )

        tytul.setStyleSheet(
            """
            font-size: 24px;
            font-weight: bold;
            padding: 10px;
            """
        )

        layout.addWidget(
            tytul
        )


        # ======================================
        # FILTRY
        # ======================================

        grupa_filtry = QGroupBox(
            "Filtry"
        )

        filtry_layout = QVBoxLayout()


        # ======================================
        # FILTR LEKÓW
        # ======================================

        wiersz_leki = QHBoxLayout()

        wiersz_leki.addWidget(
            QLabel("Leki:")
        )

        self.filtr = QComboBox()

        self.filtr.addItem(
            "Wszyscy z gorączką"
        )

        self.filtr.addItem(
            "Tylko biorący leki"
        )

        self.filtr.addItem(
            "Tylko niebiorący leków"
        )


        # Pobranie nazw leków
        leki = goraczka[
            goraczka["Czy brał(a) leki"] == "Tak"
        ]["Jakie leki"].dropna().unique()


        for lek in leki:

            self.filtr.addItem(
                str(lek)
            )


        wiersz_leki.addWidget(
            self.filtr
        )

        filtry_layout.addLayout(
            wiersz_leki
        )


        # ======================================
        # SUWAK WIEKU
        # ======================================

        self.label_wiek = QLabel(
            "Wiek: 0 - 100 lat"
        )

        self.suwak_wiek_min = QSlider(
            Qt.Orientation.Horizontal
        )

        self.suwak_wiek_min.setMinimum(
            0
        )

        self.suwak_wiek_min.setMaximum(
            100
        )

        self.suwak_wiek_min.setValue(
            0
        )


        self.suwak_wiek_max = QSlider(
            Qt.Orientation.Horizontal
        )

        self.suwak_wiek_max.setMinimum(
            0
        )

        self.suwak_wiek_max.setMaximum(
            100
        )

        self.suwak_wiek_max.setValue(
            100
        )


        self.suwak_wiek_min.valueChanged.connect(
            self.zmien_etykiete
        )

        self.suwak_wiek_max.valueChanged.connect(
            self.zmien_etykiete
        )


        filtry_layout.addWidget(
            self.label_wiek
        )

        filtry_layout.addWidget(
            self.suwak_wiek_min
        )

        filtry_layout.addWidget(
            self.suwak_wiek_max
        )


        # ======================================
        # SUWAK BMI
        # ======================================

        self.label_bmi = QLabel(
            "BMI: 15.0 - 40.0"
        )

        self.suwak_bmi_min = QSlider(
            Qt.Orientation.Horizontal
        )

        self.suwak_bmi_min.setMinimum(
            150
        )

        self.suwak_bmi_min.setMaximum(
            400
        )

        self.suwak_bmi_min.setValue(
            150
        )


        self.suwak_bmi_max = QSlider(
            Qt.Orientation.Horizontal
        )

        self.suwak_bmi_max.setMinimum(
            150
        )

        self.suwak_bmi_max.setMaximum(
            400
        )

        self.suwak_bmi_max.setValue(
            400
        )


        self.suwak_bmi_min.valueChanged.connect(
            self.zmien_etykiete
        )

        self.suwak_bmi_max.valueChanged.connect(
            self.zmien_etykiete
        )


        filtry_layout.addWidget(
            self.label_bmi
        )

        filtry_layout.addWidget(
            self.suwak_bmi_min
        )

        filtry_layout.addWidget(
            self.suwak_bmi_max
        )


        # ======================================
        # SUWAK TEMPERATURY
        # ======================================

        self.label_temp = QLabel(
            "Temperatura: 36.0 - 42.0 °C"
        )

        self.suwak_temp_min = QSlider(
            Qt.Orientation.Horizontal
        )

        self.suwak_temp_min.setMinimum(
            360
        )

        self.suwak_temp_min.setMaximum(
            420
        )

        self.suwak_temp_min.setValue(
            360
        )


        self.suwak_temp_max = QSlider(
            Qt.Orientation.Horizontal
        )

        self.suwak_temp_max.setMinimum(
            360
        )

        self.suwak_temp_max.setMaximum(
            420
        )

        self.suwak_temp_max.setValue(
            420
        )


        self.suwak_temp_min.valueChanged.connect(
            self.zmien_etykiete
        )

        self.suwak_temp_max.valueChanged.connect(
            self.zmien_etykiete
        )


        filtry_layout.addWidget(
            self.label_temp
        )

        filtry_layout.addWidget(
            self.suwak_temp_min
        )

        filtry_layout.addWidget(
            self.suwak_temp_max
        )


        # ======================================
        # PRZYCISKI
        # ======================================

        przyciski = QHBoxLayout()


        # --------------------------------------
        # POKAŻ DANE
        # --------------------------------------

        self.przycisk = QPushButton(
            "🔎 Pokaż dane"
        )

        self.przycisk.clicked.connect(
            self.aktualizuj
        )

        przyciski.addWidget(
            self.przycisk
        )


        # --------------------------------------
        # ZAPISZ DO EXCELA
        # --------------------------------------

        self.przycisk_excel = QPushButton(
            "📊 Zapisz wyniki do Excel"
        )

        self.przycisk_excel.clicked.connect(
            self.eksport_do_excel
        )

        przyciski.addWidget(
            self.przycisk_excel
        )
        print("✅ EKSPORT DZIAŁA!")
        print ("Git działa")
        filtry_layout.addLayout(
            przyciski
        )


        grupa_filtry.setLayout(
            filtry_layout
        )

        layout.addWidget(
            grupa_filtry
        )


        # ======================================
        # INFORMACJA O LICZBIE PACJENTÓW
        # ======================================

        self.informacja = QLabel(
            "Znaleziono pacjentów: 0"
        )

        self.informacja.setStyleSheet(
            """
            font-size: 16px;
            font-weight: bold;
            padding: 5px;
            """
        )

        layout.addWidget(
            self.informacja
        )


        # ======================================
        # WYKRES
        # ======================================

        self.figure = Figure(
            figsize=(10, 5)
        )

        self.canvas = FigureCanvas(
            self.figure
        )

        layout.addWidget(
            self.canvas
        )


        # ======================================
        # TABELA
        # ======================================

        self.tabela = QTableWidget()

        self.tabela.setColumnCount(
            8
        )

        self.tabela.setHorizontalHeaderLabels([
            "Imię",
            "Płeć",
            "Wiek",
            "Temperatura",
            "Leki",
            "Jakie leki",
            "Ciśnienie",
            "BMI"
        ])

        layout.addWidget(
            self.tabela
        )


        self.setLayout(
            layout
        )


        # ======================================
        # PIERWSZE WYŚWIETLENIE
        # ======================================

        self.aktualizuj()


    # ==========================================
    # AKTUALIZACJA NAPISÓW SUWAKÓW
    # ==========================================

    def zmien_etykiete(self):

        wiek_min = (
            self.suwak_wiek_min.value()
        )

        wiek_max = (
            self.suwak_wiek_max.value()
        )

        bmi_min = (
            self.suwak_bmi_min.value() / 10
        )

        bmi_max = (
            self.suwak_bmi_max.value() / 10
        )

        temp_min = (
            self.suwak_temp_min.value() / 10
        )

        temp_max = (
            self.suwak_temp_max.value() / 10
        )


        self.label_wiek.setText(
            f"Wiek: {wiek_min} - "
            f"{wiek_max} lat"
        )

        self.label_bmi.setText(
            f"BMI: {bmi_min:.1f} - "
            f"{bmi_max:.1f}"
        )

        self.label_temp.setText(
            f"Temperatura: "
            f"{temp_min:.1f} - "
            f"{temp_max:.1f} °C"
        )


    # ==========================================
    # POBIERANIE PRZEFILTROWANYCH DANYCH
    # ==========================================

    def pobierz_przefiltrowane_dane(self):

        wybor = self.filtr.currentText()

        dane_gui = goraczka.copy()


        # ======================================
        # FILTR LEKÓW
        # ======================================

        if wybor == "Tylko biorący leki":

            dane_gui = dane_gui[
                dane_gui["Czy brał(a) leki"] == "Tak"
            ]


        elif wybor == "Tylko niebiorący leków":

            dane_gui = dane_gui[
                dane_gui["Czy brał(a) leki"] == "Nie"
            ]


        elif wybor != "Wszyscy z gorączką":

            dane_gui = dane_gui[
                dane_gui["Jakie leki"]
                .astype(str)
                .str.contains(
                    wybor,
                    case=False,
                    na=False
                )
            ]


        # ======================================
        # FILTR WIEKU
        # ======================================

        wiek_min = (
            self.suwak_wiek_min.value()
        )

        wiek_max = (
            self.suwak_wiek_max.value()
        )


        if wiek_min > wiek_max:

            wiek_min, wiek_max = (
                wiek_max,
                wiek_min
            )


        dane_gui = dane_gui[
            (dane_gui["Wiek"] >= wiek_min) &
            (dane_gui["Wiek"] <= wiek_max)
        ]


        # ======================================
        # FILTR BMI
        # ======================================

        bmi_min = (
            self.suwak_bmi_min.value() / 10
        )

        bmi_max = (
            self.suwak_bmi_max.value() / 10
        )


        if bmi_min > bmi_max:

            bmi_min, bmi_max = (
                bmi_max,
                bmi_min
            )


        dane_gui = dane_gui[
            (dane_gui["BMI"] >= bmi_min) &
            (dane_gui["BMI"] <= bmi_max)
        ]


        # ======================================
        # FILTR TEMPERATURY
        # ======================================

        temp_min = (
            self.suwak_temp_min.value() / 10
        )

        temp_max = (
            self.suwak_temp_max.value() / 10
        )


        if temp_min > temp_max:

            temp_min, temp_max = (
                temp_max,
                temp_min
            )


        dane_gui = dane_gui[
            (dane_gui["Temperatura (°C)"] >= temp_min) &
            (dane_gui["Temperatura (°C)"] <= temp_max)
        ]


        return dane_gui


    # ==========================================
    # AKTUALIZACJA DANYCH
    # ==========================================

    def aktualizuj(self):

        dane_gui = (
            self.pobierz_przefiltrowane_dane()
        )


        # ======================================
        # LICZBA PACJENTÓW
        # ======================================

        self.informacja.setText(
            f"Znaleziono pacjentów: "
            f"{len(dane_gui)} / "
            f"{len(goraczka)}"
        )


        # ======================================
        # TABELA
        # ======================================

        self.tabela.setRowCount(
            len(dane_gui)
        )


        for wiersz, (_, pacjent) in enumerate(
            dane_gui.iterrows()
        ):

            wartosci = [
                pacjent["Imię (dane fikcyjne)"],
                pacjent["Płeć"],
                pacjent["Wiek"],
                pacjent["Temperatura (°C)"],
                pacjent["Czy brał(a) leki"],
                pacjent["Jakie leki"],
                pacjent["Ciśnienie (mmHg)"],
                pacjent["BMI"]
            ]


            for kolumna, wartosc in enumerate(
                wartosci
            ):

                if pd.isna(wartosc):

                    wartosc = ""


                self.tabela.setItem(
                    wiersz,
                    kolumna,
                    QTableWidgetItem(
                        str(wartosc)
                    )
                )


        # ======================================
        # WYKRES
        # ======================================

        self.figure.clear()

        ax = self.figure.add_subplot(
            111
        )


        if len(dane_gui) > 0:

            wynik_gui = pd.crosstab(
                [
                    dane_gui["Grupa wiekowa"],
                    dane_gui["Płeć"]
                ],
                dane_gui["Czy brał(a) leki"]
            )


            wynik_gui.plot(
                kind="bar",
                ax=ax
            )


            ax.set_title(
                "Pacjenci z gorączką według wieku,\n"
                "płci i przyjmowania leków"
            )

            ax.set_xlabel(
                "Grupa wiekowa i płeć"
            )

            ax.set_ylabel(
                "Liczba pacjentów"
            )

            ax.tick_params(
                axis="x",
                rotation=45
            )

            ax.legend(
                title="Czy brał(a) leki"
            )


            # ==================================
            # LICZBY NA SŁUPKACH
            # ==================================

            for kontener in ax.containers:

                ax.bar_label(
                    kontener,
                    fmt="%d",
                    padding=3
                )


        else:

            ax.text(
                0.5,
                0.5,
                "Brak pacjentów spełniających kryteria",
                ha="center",
                va="center",
                fontsize=14
            )

            ax.set_axis_off()


        self.figure.tight_layout()

        self.canvas.draw()


    # ==========================================
    # ZAPIS WYNIKÓW DO NOWEGO EXCELA
    # ==========================================

    def eksport_do_excel(self):

        # Pobieramy dokładnie te same dane,
        # które są aktualnie widoczne w GUI
        dane_excel = (
            self.pobierz_przefiltrowane_dane()
            .copy()
        )


        try:

            # ======================================
            # UTWORZENIE NOWEGO PLIKU
            # ======================================

            book = Workbook()


            # Pierwszy arkusz
            arkusz = book.active

            arkusz.title = (
                "Wyniki GUI"
            )


            # ======================================
            # NAGŁÓWKI
            # ======================================

            for kolumna, nazwa in enumerate(
                dane_excel.columns,
                start=1
            ):

                arkusz.cell(
                    row=1,
                    column=kolumna,
                    value=nazwa
                )


            # ======================================
            # DANE
            # ======================================

            for wiersz, dane_wiersza in enumerate(
                dane_excel.itertuples(
                    index=False
                ),
                start=2
            ):

                for kolumna, wartosc in enumerate(
                    dane_wiersza,
                    start=1
                ):

                    # Puste wartości
                    if pd.isna(wartosc):

                        wartosc = None


                    # Konwersja typów numpy
                    if hasattr(
                        wartosc,
                        "item"
                    ):

                        try:

                            wartosc = (
                                wartosc.item()
                            )

                        except Exception:

                            pass


                    # Nietypowe wartości
                    # zamieniamy na tekst
                    if not isinstance(
                        wartosc,
                        (
                            str,
                            int,
                            float,
                            bool,
                            type(None)
                        )
                    ):

                        wartosc = str(
                            wartosc
                        )


                    arkusz.cell(
                        row=wiersz,
                        column=kolumna,
                        value=wartosc
                    )


            # ======================================
            # SZEROKOŚĆ KOLUMN
            # ======================================

            for kolumna in arkusz.columns:

                maksymalna_dlugosc = 0

                litera = (
                    kolumna[0].column_letter
                )


                for komorka in kolumna:

                    if komorka.value is not None:

                        dlugosc = len(
                            str(
                                komorka.value
                            )
                        )

                        maksymalna_dlugosc = max(
                            maksymalna_dlugosc,
                            dlugosc
                        )


                arkusz.column_dimensions[
                    litera
                ].width = min(
                    maksymalna_dlugosc + 2,
                    35
                )


            # ======================================
            # ZAMROŻENIE NAGŁÓWKA
            # ======================================

            arkusz.freeze_panes = "A2"


            # ======================================
            # ZAPIS NOWEGO PLIKU
            # ======================================

            book.save(
                SCIEZKA_WYNIKI
            )


            # ======================================
            # AUTOMATYCZNE OTWARCIE EXCELA
            # ======================================

            if sys.platform.startswith("win"):

                os.startfile(
                    SCIEZKA_WYNIKI
                )

            elif sys.platform == "darwin":

                subprocess.Popen([
                    "open",
                    SCIEZKA_WYNIKI
                ])

            else:

                subprocess.Popen([
                    "xdg-open",
                    SCIEZKA_WYNIKI
                ])


            # ======================================
            # KOMUNIKAT
            # ======================================

            QMessageBox.information(
                self,
                "Zapisano",
                "Utworzono nowy plik Excel!\n\n"
                "Nazwa pliku:\n"
                "wyniki_gui.xlsx\n\n"
                f"Liczba pacjentów: "
                f"{len(dane_excel)}"
            )


        except PermissionError:

            QMessageBox.warning(
                self,
                "Błąd zapisu",
                "Nie można zapisać pliku.\n\n"
                "Plik 'wyniki_gui.xlsx' może być "
                "aktualnie otwarty w Excelu.\n\n"
                "Zamknij go i spróbuj ponownie."
            )


        except Exception as e:

            QMessageBox.critical(
                self,
                "Błąd",
                "Wystąpił problem podczas "
                "tworzenia pliku Excel:\n\n"
                f"{e}"
            )


# ==========================================
# URUCHOMIENIE PROGRAMU
# ==========================================

app = QApplication(
    sys.argv
)

okno = Okno()

okno.show()

sys.exit(
    app.exec()
)